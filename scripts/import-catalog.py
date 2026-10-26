"""Import source XLS and supplementary XLSX without changing the source workbooks.
Usage: python scripts/import-catalog.py full.xls selected.xlsx
Dependencies: xlrd, openpyxl. Run from the repository root.
"""
import collections
import hashlib
import json
from pathlib import Path
import re
import sys

import xlrd
from openpyxl import load_workbook


def text(value):
    if value is None:
        return ''
    if isinstance(value, (float, int)) and value == int(value):
        return str(int(value))
    return re.sub(r'\s+', ' ', str(value)).strip()


CATEGORIES = [
    ('filters', 'Фильтры'), ('oils', 'Масла и технические жидкости'),
    ('engine', 'Детали двигателя'), ('hydraulics', 'Гидравлика'),
    ('undercarriage', 'Ходовая часть'), ('electrics', 'Электрика'),
    ('transmission', 'Трансмиссия'), ('other', 'Другие запчасти и материалы'),
]


def category(name):
    name = name.casefold().replace('ё', 'е')
    # Conservative grouping by the item name. Unclear items remain in "other".
    for key, pattern in [
        ('filters', r'\bфильтр|\bфильтроэлемент'),
        ('oils', r'^(масло\b|смазка\b|антифриз\b|тосол\b|жидкость\b)'),
        ('electrics', r'\b(датчик|генератор|стартер|реле|предохранитель|катушка|выключатель|аккумулятор|фара|лампа|электродвигатель|проводка|электроклапан)'),
        ('hydraulics', r'\b(гидро|рвд\b|распределитель)'),
        ('engine', r'\b(поршень|поршнев|коленвал|коленчат|распредвал|двигател|форсун|турбокомпрессор|турбина|вкладыш|гильза|тнвд\b)'),
        ('undercarriage', r'\b(гусениц|гусенич|бандаж|опорный каток|поддерживающий каток|башмак|трак\b|шина\b|ступица|балансир)'),
        ('transmission', r'\b(кпп\b|трансмисси|сцеплени|кардан|полуось|дифференциал|коробка передач)'),
    ]:
        if re.search(pattern, name):
            return key
    return 'other'


def main():
    full_path, selected_path = map(Path, sys.argv[1:3])
    sheet = xlrd.open_workbook(full_path).sheet_by_index(0)
    products, skipped = [], []
    for row_number in range(1, sheet.nrows):
        code, article, maker, name, _ = map(text, sheet.row_values(row_number))
        # Empty article and manufacturer identify nine directory/group rows.
        if not name or not article and not maker:
            skipped.append({'row': row_number + 1, 'name': name, 'reason': 'directory heading'})
            continue
        if article == '0':
            article = ''  # Zero in the source is not a usable part number.
        products.append(dict(id=code, name=name, article=article, manufacturer=maker,
                             category=category(name)))

    def key(article, maker):
        return (re.sub(r'[^\w]', '', article).casefold(), maker.casefold())

    known = {key(p['article'], p['manufacturer']) for p in products if p['article']}
    supplement = load_workbook(selected_path, data_only=True)
    supplemental, overlaps = 0, 0
    for row in supplement.worksheets[0].iter_rows(min_row=4, values_only=True):
        short_name, name, maker, article = map(text, row[1:5])
        name = name or short_name
        if not name:
            continue
        identity = key(article, maker)
        if article and identity in known:
            overlaps += 1
            continue
        digest = hashlib.sha256('|'.join([name, article, maker]).encode()).hexdigest()[:16]
        products.append(dict(id='selected-' + digest, name=name, article=article,
                             manufacturer=maker, category=category(name)))
        if article:
            known.add(identity)
        supplemental += 1

    ids = [p['id'] for p in products]
    assert len(ids) == len(set(ids)), 'Duplicate product IDs'
    selected_brands = sorted({text(row[1]) for row in supplement.worksheets[1].values
                              if row[1] and row[1] != 'Производитель' and text(row[1]) != 'Масло'})
    data = dict(categories=[dict(id=k, name=n) for k, n in CATEGORIES],
                manufacturers=sorted({p['manufacturer'] for p in products if p['manufacturer']} | set(selected_brands)),
                products=products)
    Path('assets/js/catalog-data.js').write_text(
        "'use strict';\n// Imported from supplied workbooks. No prices or stock claims.\nwindow.MTG_CATALOG = "
        + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
    report = dict(sourceRows=sheet.nrows-1, fullProducts=len(products)-supplemental,
                  supplementalProducts=supplemental, supplementaryMatches=overlaps,
                  totalProducts=len(products), manufacturerLabels=len(data['manufacturers']),
                  missingArticle=sum(not p['article'] for p in products),
                  missingManufacturer=sum(not p['manufacturer'] for p in products),
                  categories=dict(collections.Counter(p['category'] for p in products)),
                  excludedRows=skipped,
                  notes=['Source labels retained, including brands and country labels.',
                         'Categories inferred conservatively from item names; other requires further classification.',
                         'No product photos, stock levels or prices supplied.'])
    Path('catalog-import-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
